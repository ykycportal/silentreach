"""
Output formatters for SilentReach.
Supports JSON, Markdown, CSV, PDF, Excel, and ODS formats.
"""

import json
import csv
import io
import logging
from typing import List, Dict, Any
from datetime import datetime
from pathlib import Path

logger = logging.getLogger(__name__)

try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    HAS_REPORTLAB = True
except ImportError:
    # Try fpdf2 as fallback (lighter dependency)
    try:
        from fpdf import FPDF
        HAS_FPDF = True
    except ImportError:
        HAS_FPDF = False
    HAS_REPORTLAB = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import odf.opendocument
    import odf.table
    import odf.text
    HAS_ODF = True
except ImportError:
    HAS_ODF = False

logger = logging.getLogger(__name__)


class OutputFormatter:
    """Formats scraper results into various output formats."""
    
    @staticmethod
    def to_json(data: Dict, pretty: bool = True) -> str:
        """Convert data to JSON string."""
        if pretty:
            return json.dumps(data, indent=2, default=str)
        return json.dumps(data, default=str)
    
    @staticmethod
    def to_markdown(data: Dict, title: str = "Results") -> str:
        """Convert data to Markdown format."""
        lines = [
            f"# {title}",
            f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            "",
        ]
        
        if isinstance(data, dict):
            for key, value in data.items():
                if isinstance(value, list):
                    lines.append(f"## {key.capitalize()} ({len(value)} items)")
                    lines.append("")
                    
                    for i, item in enumerate(value[:50], 1):
                        if isinstance(item, dict):
                            # Get first few fields
                            preview = ", ".join(str(v) for v in list(item.values())[:3])
                            lines.append(f"{i}. {preview}")
                        else:
                            lines.append(f"{i}. {item}")
                    
                    lines.append("")
                else:
                    lines.append(f"**{key.capitalize()}**: {value}")
                    lines.append("")
        
        return "\n".join(lines)
    
    @staticmethod
    def to_csv(data: List[Dict], filename: str = None) -> str:
        """Convert data to CSV format."""
        if not data:
            return ""
        
        # Get all unique keys
        all_keys = set()
        for item in data:
            if isinstance(item, dict):
                all_keys.update(item.keys())
        
        keys = sorted(all_keys)
        
        # Write CSV
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=keys, extrasaction='ignore')
        writer.writeheader()
        
        for item in data:
            if isinstance(item, dict):
                # Convert values to strings
                row = {k: str(v) if v is not None else "" for k, v in item.items()}
                writer.writerow(row)
        
        csv_content = output.getvalue()
        output.close()
        
        # Save to file if filename provided
        if filename:
            with open(filename, "w", newline="") as f:
                f.write(csv_content)
            logger.info(f"CSV saved to {filename}")
        
        return csv_content
    
    @staticmethod
    def to_text(data: List[Dict], max_length: int = 1000) -> str:
        """Convert data to plain text."""
        texts = []
        
        for item in data:
            if isinstance(item, dict):
                # Combine relevant fields
                text_parts = []
                for key in ["text", "content", "title", "name", "description"]:
                    if key in item and item[key]:
                        text_parts.append(str(item[key]))
                
                if text_parts:
                    texts.append(" ".join(text_parts)[:max_length])
            else:
                texts.append(str(item)[:max_length])
        
        return "\n\n".join(texts)
    
    @staticmethod
    def _normalize_for_csv_txt(data: Any) -> List[Dict]:
        """Flatten nested platform data into a list of row dicts for CSV/TXT export."""
        rows = []
        if isinstance(data, dict):
            for platform, results in data.items():
                if isinstance(results, dict):
                    # Try common result field names
                    for key in ['posts', 'videos', 'tweets', 'topics', 'notes', 'profiles', 'data']:
                        items = results.get(key, [])
                        if items:
                            for item in items:
                                if isinstance(item, dict):
                                    row = {'platform': platform, **item}
                                    rows.append(row)
                            break
                    # If no recognized key, try the values directly
                    if not rows or not any(r.get('platform') == platform for r in rows):
                        for key, value in results.items():
                            if isinstance(value, list):
                                for item in value:
                                    if isinstance(item, dict):
                                        row = {'platform': platform, **item}
                                        rows.append(row)
        elif isinstance(data, list):
            rows = data
        return rows

    @staticmethod
    def save_to_file(data: Any, filename: str, format: str = "auto") -> str:
        """
        Save data to file with automatic format detection.

        Returns:
            Path to saved file or None if failed
        """
        path = Path(filename)

        # Auto-detect format from extension
        if format == "auto":
            fmt = path.suffix.lower().lstrip('.')
        else:
            fmt = format.lower()

        # Ensure proper extension
        if fmt not in ["json", "md", "csv", "txt", "pdf", "xlsx", "ods"]:
            fmt = "json"
            path = path.with_suffix(f".{fmt}")
        elif not path.suffix:
            path = path.with_suffix(f".{fmt}")

        content = None

        if fmt == "json":
            content = OutputFormatter.to_json(data)
        elif fmt in ["md", "markdown"]:
            content = OutputFormatter.to_markdown(data)
        elif fmt == "csv":
            # Flatten nested data structure for CSV export
            flat_data = OutputFormatter._normalize_for_csv_txt(data)
            content = OutputFormatter.to_csv(flat_data)
        elif fmt == "txt":
            # Flatten nested data structure for TXT export
            flat_data = OutputFormatter._normalize_for_csv_txt(data)
            content = OutputFormatter.to_text(flat_data)
        elif fmt == "pdf":
            content = OutputFormatter.to_pdf(data, return_string=True)
        elif fmt == "xlsx":
            content = OutputFormatter.to_excel(data, return_string=True)
        elif fmt == "ods":
            content = OutputFormatter.to_odf(data, return_string=True)

        if content:
            path.parent.mkdir(parents=True, exist_ok=True)
            with open(path, "w" if fmt in ["json", "md", "csv", "txt"] else "wb") as f:
                f.write(content)

            logger.info(f"Data saved to {path}")
            return str(path)

        logger.error(f"Failed to generate {fmt} content")
        return None
    
    @staticmethod
    def to_pdf(data: Dict, title: str = "SilentReach Report", filename: str = None, 
               return_string: bool = True) -> Any:
        """Convert data to PDF format (reportlab or fpdf2 fallback)."""
        try:
            # Try reportlab first
            if HAS_REPORTLAB:
                from reportlab.lib.pagesizes import A4
                from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
                from reportlab.lib import colors
                from reportlab.lib.units import inch
                from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
                
                buffer = io.BytesIO() if return_string else None
                if not return_string:
                    Path(filename).parent.mkdir(parents=True, exist_ok=True)
                
                doc = SimpleDocTemplate(
                    buffer if buffer else filename,
                    pagesize=A4,
                    rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=72
                )
                
                styles = getSampleStyleSheet()
                title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'],
                    fontSize=24, spaceAfter=30, textColor=colors.HexColor('#1a1a2e'))
                subtitle_style = ParagraphStyle('Subtitle', parent=styles['Normal'],
                    fontSize=12, textColor=colors.HexColor('#666666'), spaceAfter=20)
                
                story = []
                story.append(Paragraph(title, title_style))
                story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", subtitle_style))
                story.append(Spacer(1, 20))
                
                if isinstance(data, dict):
                    for platform, results in data.items():
                        if isinstance(results, dict):
                            story.append(Paragraph(platform.capitalize(), styles['Heading2']))
                            items = results.get('posts', results.get('videos', results.get('tweets', [])))
                            if items:
                                table_data = [['Title', 'Author', 'Score', 'Date']]
                                for item in items[:20]:
                                    table_data.append([
                                        str(item.get('title', item.get('text', '')))[:50],
                                        str(item.get('author', ''))[:20],
                                        str(item.get('score', item.get('likes', ''))),
                                        str(item.get('created_at', item.get('timestamp', '')))[:10],
                                    ])
                                tbl = Table(table_data, colWidths=[4*inch, 1.5*inch, 1*inch, 1.5*inch])
                                tbl.setStyle(TableStyle([
                                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#667eea')),
                                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                                    ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f8f9fa')),
                                    ('GRID', (0, 0), (-1, -1), 1, colors.HexColor('#e0e0e0')),
                                ]))
                                story.append(tbl)
                            else:
                                story.append(Paragraph(f"No results found for {platform}", styles['Normal']))
                        story.append(Spacer(1, 20))
                
                doc.build(story)
                if return_string:
                    buffer.seek(0)
                    return buffer.read()
                return filename
            # Fall back to fpdf2
            elif HAS_FPDF:
                from fpdf import FPDF
                pdf = FPDF()
                pdf.add_page()
                # Try to use a Unicode-capable font if available
                try:
                    pdf.add_font('DejaVu', '', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
                    pdf.set_font('DejaVu', size=11)
                except:
                    pdf.set_font('helvetica', size=11)
                pdf.set_font('helvetica', size=11)
                
                # Title
                pdf.set_font('helvetica', 'B', size=18)
                pdf.cell(0, 12, title, new_x="LMARGIN", new_y="NEXT")
                pdf.set_font('helvetica', size=9)
                pdf.ln(2)
                pdf.cell(0, 6, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", new_x="LMARGIN", new_y="NEXT")
                pdf.ln(8)
                
                if isinstance(data, dict):
                    for platform, results in data.items():
                        if isinstance(results, dict):
                            pdf.set_font('helvetica', 'B', size=12)
                            pdf.cell(0, 8, platform.capitalize(), new_x="LMARGIN", new_y="NEXT")
                            pdf.ln(4)
                            items = results.get('posts', results.get('videos', results.get('tweets', [])))
                            if items:
                                pdf.set_font('helvetica', size=9)
                                for item in items[:30]:
                                    text = str(item.get('title', item.get('text', '')))[:80]
                                    author = str(item.get('author', ''))[:20]
                                    score = str(item.get('score', item.get('likes', '')))
                                    pdf.cell(0, 5, f"  {text}  |  {author}  |  score={score}", new_x="LMARGIN", new_y="NEXT")
                                pdf.ln(6)
                            else:
                                pdf.set_font('helvetica', size=10)
                                pdf.cell(0, 6, f"  No results for {platform}", new_x="LMARGIN", new_y="NEXT")
                                pdf.ln(6)
                
                if return_string:
                    return pdf.output(dest='S')
                Path(filename).parent.mkdir(parents=True, exist_ok=True)
                pdf.output(filename)
                return filename
            else:
                logger.error("Neither reportlab nor fpdf2 is installed. Run: pip install fpdf2")
                return None
        except Exception as e:
            logger.error(f"PDF generation failed: {e}")
            return None
    
    @staticmethod
    def to_excel(data: Dict, title: str = "SilentReach Report", filename: str = None,
                 return_string: bool = True) -> Any:
        """Convert data to Excel format using openpyxl."""
        if not HAS_OPENPYXL:
            logger.error("openpyxl not installed. Run: pip install openpyxl")
            return None
        
        try:
            import openpyxl
            from openpyxl.styles import Font, Fill, PatternFill, Alignment
            
            wb = openpyxl.Workbook()
            
            # Process each platform
            if isinstance(data, dict):
                for platform, results in data.items():
                    if isinstance(results, dict):
                        # Create sheet for each platform
                        ws = wb.create_sheet(title=platform[:31])  # Excel sheet name limit
                        
                        items = results.get('posts', results.get('videos', results.get('tweets', [])))
                        
                        if items:
                            # Headers
                            headers = ['Title', 'Author', 'Score', 'Date', 'URL', 'Content']
                            ws.append(headers)
                            
                            # Style headers
                            for cell in ws[1]:
                                cell.font = Font(bold=True, color='FFFFFF')
                                cell.fill = PatternFill(start_color='667eea', end_color='667eea', fill_type='solid')
                                cell.alignment = Alignment(horizontal='center')
                            
                            # Data rows
                            for item in items[:100]:  # Limit to 100 rows
                                row = [
                                    item.get('title', item.get('text', '')),
                                    item.get('author', ''),
                                    item.get('score', item.get('likes', '')),
                                    item.get('created_at', item.get('timestamp', '')),
                                    item.get('link', item.get('url', '')),
                                    str(item.get('text', item.get('description', '')))[:500],
                                ]
                                ws.append(row)
                            
                            # Auto-adjust column widths
                            for column in ws.columns:
                                max_length = 0
                                column_letter = column[0].column_letter
                                for cell in column:
                                    try:
                                        if len(str(cell.value)) > max_length:
                                            max_length = len(str(cell.value))
                                    except:
                                        pass
                                adjusted_width = min(max_length + 2, 50)
                                ws.column_dimensions[column_letter].width = adjusted_width
                        
                        # Freeze panes
                        ws.freeze_panes = 'A2'
            
            # Save
            if return_string:
                buffer = io.BytesIO()
                wb.save(buffer)
                buffer.seek(0)
                return buffer.read()
            else:
                Path(filename).parent.mkdir(parents=True, exist_ok=True)
                wb.save(filename)
                return filename
            
        except Exception as e:
            logger.error(f"Excel generation failed: {e}")
            return None
    
    @staticmethod
    def to_odf(data: Dict, title: str = "SilentReach Report", filename: str = None,
               return_string: bool = True) -> Any:
        """Convert data to OpenDocument Format (ODS) using odfpy."""
        if not HAS_ODF:
            logger.error("odfpy not installed. Run: pip install odfpy")
            return None

        try:
            from odf.opendocument import OpenDocumentSpreadsheet
            from odf.table import Table, TableRow, TableCell
            from odf.text import P

            doc = OpenDocumentSpreadsheet()

            # Process each platform
            if isinstance(data, dict):
                for platform, results in data.items():
                    if isinstance(results, dict):
                        table = Table(name=platform[:31])

                        items = results.get('posts', results.get('videos', results.get('tweets', [])))

                        if items:
                            # Header row
                            header_row = TableRow()
                            for header in ['Title', 'Author', 'Score', 'Date', 'URL']:
                                cell = TableCell()
                                p = P(text=header)
                                cell.addElement(p)
                                header_row.addElement(cell)
                            table.addElement(header_row)

                            # Data rows
                            for item in items[:100]:
                                row = TableRow()
                                for value in [
                                    item.get('title', item.get('text', '')),
                                    item.get('author', ''),
                                    item.get('score', item.get('likes', '')),
                                    item.get('created_at', item.get('timestamp', '')),
                                    item.get('link', item.get('url', '')),
                                ]:
                                    cell = TableCell()
                                    p = P(text=str(value)[:100])
                                    cell.addElement(p)
                                    row.addElement(cell)
                                table.addElement(row)

                        doc.spreadsheet.addElement(table)

            # Save
            if return_string:
                buffer = io.BytesIO()
                doc.save(buffer)
                buffer.seek(0)
                return buffer.read()
            else:
                Path(filename).parent.mkdir(parents=True, exist_ok=True)
                doc.save(filename)
                return filename

        except Exception as e:
            logger.error(f"ODF generation failed: {e}")
            return None
